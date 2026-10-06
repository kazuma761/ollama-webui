"""Reads a Word file, finds the places meant to be filled, and writes into them.

Nothing here talks to a model, and nothing here builds a Word file from
scratch. Every change is made inside the file that was uploaded, to the one
paragraph or run it concerns, so the rest - styles, fonts, headers, footers,
logos, margins, table layout - is exactly what it was.

The pieces:

- `WordFile` opens a .docx or .dotx and numbers every paragraph (`p1`, `p2`,
  ... in the body, `h1`.. in headers, `f1`.. in footers) and every table.
- `WordFile.fields()` lists the blanks: marked placeholders such as
  `{{name}}` or `[Client Name]`, underscore and dotted blanks, highlighted
  text, empty table cells beside a label, "Label:" lines, empty sections
  under a heading, and content controls still showing their prompt.
- `WordFile.fill()` writes one value per blank. Each value takes the look of
  what it replaces; a label that shared a run with its blank stays as it was.

Editing by instruction lives in `wordedit.py`, which builds on this file.
"""

from __future__ import annotations

import copy
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

import docx
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ollama_pipeline import UnsupportedDocument

MAX_FIELDS = 200

_TEMPLATE_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
_DOCUMENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # old .doc, or a password-protected .docx

W_P, W_R, W_TBL, W_TR, W_TC = qn("w:p"), qn("w:r"), qn("w:tbl"), qn("w:tr"), qn("w:tc")
W_PPR, W_RPR, W_SDT, W_TXBX = qn("w:pPr"), qn("w:rPr"), qn("w:sdt"), qn("w:txbxContent")
W_VAL = qn("w:val")
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
MC_ALTERNATE, MC_CHOICE, MC_FALLBACK = _MC + "AlternateContent", _MC + "Choice", _MC + "Fallback"

# A run made only of these can have its text replaced without losing anything else.
_TEXT_ONLY = {qn(f"w:{tag}") for tag in ("rPr", "t", "tab", "br", "cr", "noBreakHyphen", "softHyphen", "lastRenderedPageBreak")}
# Paragraph-mark formatting that is not valid on a run.
_NOT_FOR_RUNS = {qn(f"w:{tag}") for tag in ("ins", "del", "moveFrom", "moveTo", "rPrChange")}
# What a copied paragraph must not carry along: section breaks and anchors that have to stay unique.
_NOT_FOR_COPIES = {qn(f"w:{tag}") for tag in ("bookmarkStart", "bookmarkEnd", "commentRangeStart", "commentRangeEnd", "permStart", "permEnd")}

# Marked placeholders, most specific first; a later pattern never matches inside an earlier one.
_MARKED = [
    ("placeholder", re.compile(r"\{\{\s*([^{}\n]{1,80}?)\s*\}\}")),
    ("placeholder", re.compile(r"\[\[\s*([^\[\]\n]{1,80}?)\s*\]\]")),
    ("placeholder", re.compile(r"«\s*([^«»\n]{1,80}?)\s*»")),
    ("placeholder", re.compile(r"<<\s*([^<>\n]{1,80}?)\s*>>")),
    ("placeholder", re.compile(r"\[([^\[\]\n]{2,60})\]")),
    ("placeholder", re.compile(r"<([A-Za-z][\w '/&.-]{1,40})>")),
]
_BLANK = re.compile(r"_{3,}|\.{5,}|…{2,}")
# Dummy values a template leaves for its user to overwrite.
_SAMPLE = re.compile(
    r"\b[Dd]{2}[/.\-][Mm]{2}[/.\-][Yy]{2,4}\b|\b[Xx]{4,}\b|"
    r"(?i:click or tap here to enter (?:text|a date)\.?|click here to enter (?:text|a date)\.?|"
    r"(?:enter|insert|type) (?:your |the )?(?:text|[\w ]{2,30}?) here\.?)"
)
_CITATION = re.compile(r"^[\d\s,;–-]+$")
_HEADING = re.compile(r"^heading (\d)$", re.IGNORECASE)
_BULLET = re.compile(r"^\s*[-*•]\s+")


@dataclass
class Block:
    """One paragraph of the file."""

    id: str
    element: Any  # the w:p element
    where: str  # "body" | "header" | "footer" | "textbox"
    cell: tuple[int, int, int] | None = None  # (table, row, column), counted from 1
    anchor: str | None = None  # for a text box paragraph: the id of the paragraph the box hangs on


@dataclass
class Spot:
    """One place a value is written."""

    block: str
    mode: str  # "span" | "paragraph" | "append" | "section" | "after"
    start: int = 0
    end: int = 0


@dataclass
class Field:
    id: str
    label: str
    kind: str  # "text" for a value, "section" for one or more paragraphs
    how: str  # how it was found: placeholder | blank | highlight | control | sample | cell | label | section
    context: str  # the surrounding text, to tell the reader and the model what belongs here
    current: str = ""  # what stands there now
    options: list[str] = field(default_factory=list)  # choices of a drop-down control
    spots: list[Spot] = field(default_factory=list)

    def public(self) -> dict[str, Any]:
        return {
            "id": self.id, "label": self.label, "kind": self.kind, "how": self.how,
            "context": self.context, "current": self.current, "options": self.options, "places": len(self.spots),
        }


# ── Opening ──


def as_docx(name: str, data: bytes) -> bytes:
    """The file as plain .docx bytes. A .dotx template only needs relabelling; anything unusable is refused."""
    if data.startswith(_OLE_MAGIC):
        raise UnsupportedDocument(
            f"{name} is an old .doc file or is password-protected. Save it as .docx without a password and try again."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            types = archive.read("[Content_Types].xml").decode("utf-8", "replace")
            if "macroEnabled" in types:
                raise UnsupportedDocument(f"{name} contains macros. Save it as a plain .docx and try again.")
            if "wordprocessingml" not in types:
                raise UnsupportedDocument(f"{name} is not a Word document.")
            if _TEMPLATE_TYPE not in types:
                return data
            out = io.BytesIO()
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
                for item in archive.infolist():
                    content = archive.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        content = types.replace(_TEMPLATE_TYPE, _DOCUMENT_TYPE).encode("utf-8")
                    target.writestr(item, content)
            return out.getvalue()
    except (zipfile.BadZipFile, KeyError) as exc:
        raise UnsupportedDocument(f"{name} is not a Word .docx or .dotx file.") from exc


def open_word(name: str, data: bytes) -> DocumentType:
    converted = as_docx(name, data)
    try:
        return docx.Document(io.BytesIO(converted))
    except Exception as exc:
        raise UnsupportedDocument(f"{name} could not be read as a Word document.") from exc


# ── Runs and paragraphs ──


def _inside(element: Any, stop: Any, tags: tuple[str, ...]) -> bool:
    """True when `element` sits inside one of `tags` somewhere below `stop`."""
    for ancestor in element.iterancestors():
        if ancestor is stop:
            return False
        if ancestor.tag in tags:
            return True
    return False


def text_runs(paragraph: Any) -> list[Any]:
    """The runs of a paragraph whose text can be replaced safely, in reading order."""
    runs = []
    for run in paragraph.iter(W_R):
        if _inside(run, paragraph, (W_TXBX, qn("w:del"), qn("w:moveFrom"))):
            continue  # a text box anchored here, or text deleted with tracked changes
        if any(child.tag not in _TEXT_ONLY for child in run):
            continue  # holds a picture, a field code, a footnote mark ...
        if any(child.tag == qn("w:br") and child.get(qn("w:type")) in ("page", "column") for child in run):
            continue  # a page break must survive
        runs.append(run)
    return runs


def box_paragraphs(paragraph: Any) -> list[Any]:
    """The paragraphs of the text boxes hanging on a paragraph: the modern copy of each, in order."""
    found = []
    for box in paragraph.iter(W_TXBX):
        if _inside(box, paragraph, (W_TXBX, MC_FALLBACK)):
            continue
        found.extend(inner for inner in box.iter(W_P) if not _inside(inner, box, (W_TXBX,)))
    return found


def spans(paragraph: Any) -> tuple[str, list[tuple[Any, int, int]]]:
    """A paragraph's text and, for each run, the range of that text it holds."""
    parts, position = [], 0
    for run in text_runs(paragraph):
        length = len(run.text)
        parts.append((run, position, position + length))
        position += length
    return "".join(run.text for run, _, _ in parts), parts


def paragraph_text(paragraph: Any) -> str:
    return spans(paragraph)[0]


def _is_on(properties: Any, tag: str) -> bool:
    if properties is None:
        return False
    found = properties.find(qn(f"w:{tag}"))
    return found is not None and found.get(W_VAL) not in ("0", "false", "none")


def _drop(properties: Any, *tags: str) -> None:
    if properties is not None:
        for tag in tags:
            for found in properties.findall(qn(f"w:{tag}")):
                properties.remove(found)


def _value_run(like: Any, text: str, *, plain: bool) -> Any:
    """A new run that looks like `like`. With `plain`, a label's bold and capitals are not carried over."""
    run = copy.deepcopy(like)
    if plain:
        _drop(run.find(W_RPR), "b", "bCs", "caps", "smallCaps")
    run.text = text
    return run


def _as_value(run: Any) -> None:
    """Removes what marked a run as a blank: its highlight and a content control's grey prompt style."""
    properties = run.find(W_RPR)
    _drop(properties, "highlight")
    if properties is not None:
        style = properties.find(qn("w:rStyle"))
        if style is not None and style.get(W_VAL) == "PlaceholderText":
            properties.remove(style)
    for ancestor in run.iterancestors(W_SDT):
        settings = ancestor.find(qn("w:sdtPr"))
        if settings is not None:
            for flag in settings.findall(qn("w:showingPlcHdr")):
                settings.remove(flag)


def replace_span(paragraph: Any, start: int, end: int, text: str, *, plain: bool = False) -> bool:
    """Replaces characters [start, end) of a paragraph. Returns True when runs had to be merged.

    Inside one run nothing else is touched. Across several runs the value takes
    the first run's formatting and the text on either side keeps its own.
    """
    _, parts = spans(paragraph)
    touched = [(run, a, b) for run, a, b in parts if a < end and b > start]
    if not touched:
        return False
    first, first_start, _ = touched[0]
    last, last_start, _ = touched[-1]
    before = first.text[: start - first_start]
    after = last.text[end - last_start:]
    shares_a_label = bool(before.strip() or (first is last and after.strip()))

    if plain and shares_a_label:
        # The blank sits in the same run as its label ("Name: ____" typed in one go, in bold):
        # the value gets a run of its own so the label keeps its look and the value does not copy it.
        value = _value_run(first, text, plain=True)
        tail = copy.deepcopy(first) if first is last and after else None
        first.text = before
        first.addnext(value)
        if tail is not None:
            tail.text = after
            value.addnext(tail)
        if not before and first is last:
            first.getparent().remove(first)
    elif first is last:
        first.text = before + text + after
        value = first
    else:
        first.text = before + text
        value = first

    if first is not last:
        for run, _, _ in touched[1:-1]:
            run.getparent().remove(run)
        if after:
            last.text = after
        else:
            last.getparent().remove(last)
    _as_value(value)
    return first is not last


def _mark_properties(paragraph: Any) -> Any:
    """The formatting Word itself would give text typed into this paragraph, as run properties."""
    properties = paragraph.find(W_PPR)
    mark = properties.find(W_RPR) if properties is not None else None
    if mark is None:
        return None
    mark = copy.deepcopy(mark)
    for child in list(mark):
        if child.tag in _NOT_FOR_RUNS:
            mark.remove(child)
    return mark


def write_paragraph(paragraph: Any, text: str) -> None:
    """Sets the whole text of a paragraph, keeping the formatting of its first run."""
    runs = text_runs(paragraph)
    if runs:
        runs[0].text = text
        _as_value(runs[0])
        for run in runs[1:]:
            run.getparent().remove(run)
        return
    run = OxmlElement("w:r")
    mark = _mark_properties(paragraph)
    if mark is not None:
        run.append(mark)
    paragraph.append(run)
    run.text = text


def append_text(paragraph: Any, text: str) -> None:
    """Adds a value after a label ("Date:"), in the label's font but not its bold."""
    runs = text_runs(paragraph)
    if not runs:
        write_paragraph(paragraph, text)
        return
    gap = "" if paragraph_text(paragraph).endswith((" ", "\t")) else " "
    runs[-1].addnext(_value_run(runs[-1], gap + text, plain=True))


def clean_copy(paragraph: Any) -> Any:
    """A copy of a paragraph that is safe to insert elsewhere in the file."""
    clone = copy.deepcopy(paragraph)
    properties = clone.find(W_PPR)
    if properties is not None:
        for section_break in properties.findall(qn("w:sectPr")):
            properties.remove(section_break)
    for child in list(clone.iter()):
        if child.tag in _NOT_FOR_COPIES:
            child.getparent().remove(child)
    for attribute in list(clone.attrib):
        if attribute.endswith("}paraId") or attribute.endswith("}textId"):
            del clone.attrib[attribute]
    return clone


def empty_copy(paragraph: Any, *, neutral: bool = False) -> Any:
    """A copy of a paragraph with its formatting and no content.

    With `neutral`, new text gets the paragraph's font, size and colour but not
    the emphasis its first words happened to have (bold, italic, underline ...).
    """
    clone = clean_copy(paragraph)
    first_run = next(iter(text_runs(clone)), None)
    kept = copy.deepcopy(first_run.find(W_RPR)) if first_run is not None and first_run.find(W_RPR) is not None else None
    if neutral:
        _drop(kept, "b", "bCs", "i", "iCs", "u", "caps", "smallCaps", "strike", "highlight", "rStyle")
    for child in list(clone):
        if child.tag != W_PPR:
            clone.remove(child)
    if kept is not None:
        run = OxmlElement("w:r")
        run.append(kept)
        clone.append(run)
    return clone


# ── The file ──


class WordFile:
    def __init__(self, name: str, data: bytes):
        self.name = name
        self.doc = open_word(name, data)
        self.blocks: list[Block] = []
        self.tables: list[list[list[Any]]] = []  # tables[t][r][c] is a w:tc element
        self._styles = {s.style_id: s.name for s in self.doc.styles if s.style_id}
        self._style_levels: dict[str, int] = {}  # custom heading styles that carry an outline level
        for style in self.doc.styles:
            level = style.element.find(f"{W_PPR}/{qn('w:outlineLvl')}")
            if style.style_id and level is not None and (level.get(W_VAL) or "9").isdigit() and int(level.get(W_VAL)) < 9:
                self._style_levels[style.style_id] = int(level.get(W_VAL)) + 1
        self._index()

    # ── Structure ──

    def _index(self) -> None:
        body = self.doc.element.body
        self._boxes = 0
        cells: dict[int, tuple[int, int, int]] = {}
        for table in body.iter(W_TBL):
            if _inside(table, body, (W_TXBX,)):
                continue
            rows = [[cell for cell in row.iterchildren(W_TC)] for row in table.iterchildren(W_TR)]
            self.tables.append(rows)
            for r, row in enumerate(rows, 1):
                for c, cell in enumerate(row, 1):
                    cells[id(cell)] = (len(self.tables), r, c)

        def add(root: Any, where: str, prefix: str, count: int) -> int:
            for paragraph in root.iter(W_P):
                if _inside(paragraph, root, (W_TXBX,)):
                    continue  # reached through the paragraph its text box hangs on, below
                count += 1
                cell = next((cells[id(a)] for a in paragraph.iterancestors(W_TC) if id(a) in cells), None)
                anchor = Block(f"{prefix}{count}", paragraph, where, cell)
                self.blocks.append(anchor)
                # Word stores each text box twice: a modern copy (mc:Choice) and one for old
                # versions (mc:Fallback). Only the modern one is numbered; save() copies changes over.
                for inner in box_paragraphs(paragraph):
                    self._boxes += 1
                    self.blocks.append(Block(f"t{self._boxes}", inner, "textbox", None, anchor.id))
            return count

        add(body, "body", "p", 0)
        seen: list[Any] = []
        self._parts: list[Any] = []  # header and footer roots
        headers = footers = 0
        for section in self.doc.sections:
            for kind in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
                part = getattr(section, kind)
                if part.is_linked_to_previous:
                    continue  # no definition of its own; asking for one would create it
                element = part._element
                if any(element is known for known in seen):
                    continue
                seen.append(element)
                self._parts.append(element)
                if "header" in kind:
                    headers = add(element, "header", "h", headers)
                else:
                    footers = add(element, "footer", "f", footers)
        self._by_id = {block.id: block for block in self.blocks}
        self._by_element = {id(block.element): block for block in self.blocks}

    def block(self, block_id: str) -> Block:
        try:
            return self._by_id[block_id]
        except KeyError:
            raise KeyError(f"there is no paragraph {block_id}") from None

    def style_name(self, paragraph: Any) -> str:
        properties = paragraph.find(W_PPR)
        style = properties.find(qn("w:pStyle")) if properties is not None else None
        return self._styles.get(style.get(W_VAL), style.get(W_VAL)) if style is not None else "Normal"

    def heading_level(self, paragraph: Any) -> int | None:
        match = _HEADING.match(self.style_name(paragraph))
        if match:
            return int(match.group(1))
        properties = paragraph.find(W_PPR)
        outline = properties.find(qn("w:outlineLvl")) if properties is not None else None
        if outline is not None and (outline.get(W_VAL) or "9").isdigit() and int(outline.get(W_VAL)) < 9:
            return int(outline.get(W_VAL)) + 1
        style = properties.find(qn("w:pStyle")) if properties is not None else None
        return self._style_levels.get(style.get(W_VAL)) if style is not None else None

    def alignment(self, paragraph: Any) -> str:
        properties = paragraph.find(W_PPR)
        value = properties.find(qn("w:jc")) if properties is not None else None
        return {"center": "center", "right": "right", "end": "right", "both": "justify", "distribute": "justify"}.get(
            value.get(W_VAL) if value is not None else "", "left"
        )

    def has_style(self, name: str) -> bool:
        return any(s.name.lower() == name.lower() for s in self.doc.styles)

    def style_id(self, name: str) -> str | None:
        return next((s.style_id for s in self.doc.styles if s.name.lower() == name.lower()), None)

    @staticmethod
    def _is_empty(element: Any) -> bool:
        """No text and no picture in a paragraph or table cell."""
        paragraphs = [element] if element.tag == W_P else list(element.iter(W_P))
        if any(paragraph_text(p).strip() for p in paragraphs):
            return False
        return element.find(f".//{qn('w:drawing')}") is None and element.find(f".//{qn('w:pict')}") is None

    @staticmethod
    def _cell_text(cell: Any) -> str:
        own = (p for p in cell.iter(W_P) if not _inside(p, cell, (W_TXBX,)))
        return " ".join(t for t in (paragraph_text(p).strip() for p in own) if t)

    def is_letterhead(self) -> bool:
        """True when the design lives in headers and footers and the body is (nearly) empty.

        Only then may new text replace the body. A body with tables, pictures,
        text boxes or real text *is* the design and must be filled, not replaced.
        """
        body = self.doc.element.body
        if self.tables or body.find(f".//{qn('w:drawing')}") is not None or body.find(f".//{qn('w:pict')}") is not None:
            return False
        written = [b for b in self.blocks if b.where == "body" and paragraph_text(b.element).strip()]
        return len(written) <= 3 and sum(len(paragraph_text(b.element)) for b in written) <= 300

    def sync_text_boxes(self) -> None:
        """Copies every changed text box into its copy for old Word versions, so both say the same."""
        roots = [self.doc.element.body, *self._parts]
        for root in roots:
            for pair in root.iter(MC_ALTERNATE):
                modern = [b for b in pair.iter(W_TXBX) if _inside(b, pair, (MC_CHOICE,)) and not _inside(b, pair, (MC_FALLBACK,))]
                old = [b for b in pair.iter(W_TXBX) if _inside(b, pair, (MC_FALLBACK,))]
                if len(modern) != len(old):
                    continue
                for new, stale in zip(modern, old):
                    if [paragraph_text(p) for p in new.iter(W_P)] == [paragraph_text(p) for p in stale.iter(W_P)]:
                        continue
                    for child in list(stale):
                        stale.remove(child)
                    for child in new:
                        stale.append(copy.deepcopy(child))

    # ── Finding the blanks ──

    def _cell_label(self, position: tuple[int, int, int]) -> str:
        """What an empty table cell is for, read from the cell on its left and the top of its column."""
        t, r, c = position
        rows = self.tables[t - 1]
        row = rows[r - 1]
        left = self._cell_text(row[c - 2]) if c > 1 else ""
        first = rows[0]
        has_header = r > 1 and len(first) > 1 and all(self._cell_text(cell) for cell in first)
        header = self._cell_text(first[c - 1]) if has_header and c <= len(first) else ""
        row_name = self._cell_text(row[0]) if c > 1 else ""

        def short(text: str, limit: int) -> str:
            return text.rstrip(":").strip() if 0 < len(text) <= limit and len(text.split()) <= 8 else ""

        labelled = left.endswith(":")  # "Name:" on the left settles it, whatever the top row says
        left, header, row_name = short(left, 60), short(header, 40), short(row_name, 60)
        if header and len(first) > 2 and not labelled:
            # a grid with column headings: the row's first cell names the row
            return f"{row_name} - {header}" if row_name else f"{header} (row {r - 1})"
        return left or header

    def _blank_label(self, block: Block, text: str, start: int, end: int, taken: list[tuple[int, int]]) -> str:
        """What an underscore blank is for: the words before it, else after it, else its neighbours."""
        floor = max((b for _, b in taken if b <= start), default=0)
        # Only the words of the sentence the blank is in, after any earlier blank on the line.
        before = re.split(r"[.;!?]\s+|\n", text[floor:start].split("\t")[-1])[-1]
        before = before.strip().lstrip(".,;:)").rstrip(":-–(").strip()
        if before:
            return " ".join(before.split()[-7:])
        after = text[end:].split("\t")[0].strip().strip("()[]:").strip()
        if after and not _BLANK.match(after):
            return " ".join(after.split()[:6])
        if block.cell:
            label = self._cell_label(block.cell)
            if label:
                return label
        index = self.blocks.index(block)
        for neighbour in (index + 1, index - 1):
            if 0 <= neighbour < len(self.blocks) and self.blocks[neighbour].where == block.where:
                words = paragraph_text(self.blocks[neighbour].element).strip()
                if words and len(words) <= 60 and not _BLANK.search(words):
                    return words.rstrip(":")
        return ""

    def fields(self) -> list[Field]:
        """Every place in the file that looks meant to be filled, in reading order."""
        found: list[tuple[int, int, str | None, Field]] = []  # (block index, position, grouping key, field)

        def context_of(text: str, start: int, end: int) -> str:
            shown = text[max(0, start - 70):start] + "[____]" + text[end:end + 70]
            return " ".join(shown.split())

        for index, block in enumerate(self.blocks):
            paragraph = block.element
            text, parts = spans(paragraph)
            if not text.strip() or self.style_name(paragraph).lower().startswith("toc"):
                continue
            taken: list[tuple[int, int]] = []

            def free(a: int, b: int) -> bool:
                return all(b <= x or a >= y for x, y in taken)

            def add(how: str, label: str, start: int, end: int, key: str | None = None, options: list[str] | None = None) -> None:
                taken.append((start, end))
                spot = Spot(block.id, "span", start, end)
                item = Field("", label, "text", how, context_of(text, start, end), text[start:end], options or [], [spot])
                found.append((index, start, key, item))

            # Content controls still showing their prompt ("Click or tap here to enter text.")
            for control in paragraph.iter(W_SDT):
                settings = control.find(qn("w:sdtPr"))
                if settings is None or settings.find(qn("w:showingPlcHdr")) is None:
                    continue
                inside = [(a, b) for run, a, b in parts if any(x is control for x in run.iterancestors(W_SDT))]
                if not inside:
                    continue
                start, end = inside[0][0], inside[-1][1]
                named = settings.find(qn("w:alias")) if settings.find(qn("w:alias")) is not None else settings.find(qn("w:tag"))
                label = (named.get(W_VAL) if named is not None else "") or self._blank_label(block, text, start, end, taken)
                options = [item.get(qn("w:displayText")) or item.get(qn("w:value")) or "" for item in settings.iter(qn("w:listItem"))]
                if free(start, end):
                    add("control", label, start, end, options=[o for o in options if o])
            for control in paragraph.iterancestors(W_SDT):  # a control wrapped around the whole paragraph
                settings = control.find(qn("w:sdtPr"))
                if settings is not None and settings.find(qn("w:showingPlcHdr")) is not None and free(0, len(text)):
                    named = settings.find(qn("w:alias")) if settings.find(qn("w:alias")) is not None else settings.find(qn("w:tag"))
                    add("control", (named.get(W_VAL) if named is not None else "") or self._blank_label(block, text, 0, len(text), taken), 0, len(text))

            for how, pattern in _MARKED:
                for match in pattern.finditer(text):
                    label = " ".join(match.group(1).replace("_", " ").split())
                    if not label or _CITATION.match(label) or not free(*match.span()):
                        continue
                    label = label[0].upper() + label[1:] if label.islower() else label
                    add(how, label, match.start(), match.end(), key=label.lower())

            for match in _SAMPLE.finditer(text):
                if free(*match.span()):
                    label = self._blank_label(block, text, match.start(), match.end(), taken) or match.group(0)
                    add("sample", label, match.start(), match.end())

            for match in _BLANK.finditer(text):
                if not free(*match.span()) or re.match(r"\s*\d+\s*$", text[match.end():]):
                    continue  # dots leading to a page number are a table of contents
                add("blank", self._blank_label(block, text, match.start(), match.end(), taken), match.start(), match.end())

            # Text marked with the highlighter pen
            run_start, previous_end = None, 0
            for run, a, b in [*parts, (None, 0, 0)]:
                lit = run is not None and _is_on(run.find(W_RPR), "highlight")
                if lit and run_start is None:
                    run_start = a
                if not lit and run_start is not None:
                    end = previous_end
                    marked = text[run_start:end].strip()
                    if marked and free(run_start, end):
                        label = marked if len(marked) <= 60 else marked[:57] + "..."
                        add("highlight", label, run_start, end, key=f"highlight:{marked.lower()}")
                    run_start = None
                previous_end = b

            # "Date:" with nothing after it, unless the next table cell is the place for the value
            stripped = text.strip()
            last_in_row = not block.cell or block.cell[2] == len(self.tables[block.cell[0] - 1][block.cell[1] - 1])
            if (
                not taken and stripped.endswith(":") and len(stripped) <= 50 and len(stripped.split()) <= 6
                and self.heading_level(paragraph) is None and last_in_row
            ):
                spot = Spot(block.id, "append")
                found.append((index, len(text), None, Field("", stripped.rstrip(":").strip(), "text", "label", stripped + " [____]", "", [], [spot])))

        # Empty table cells that have a label
        used = {spot.block for _, _, _, item in found for spot in item.spots}
        for t, rows in enumerate(self.tables, 1):
            for r, row in enumerate(rows, 1):
                for c, cell in enumerate(row, 1):
                    merged = cell.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}")
                    if (merged is not None and merged.get(W_VAL) != "restart") or not self._is_empty(cell):
                        continue
                    first = next(cell.iter(W_P), None)
                    block = self._by_element.get(id(first)) if first is not None else None
                    label = self._cell_label((t, r, c))
                    if block is None or not label or block.id in used:
                        continue
                    row_text = " | ".join(self._cell_text(x) or "[____]" for x in row)
                    item = Field("", label, "text", "cell", row_text[:200], "", [], [Spot(block.id, "paragraph")])
                    found.append((self.blocks.index(block), 0, None, item))

        # A heading with nothing under it
        body = [b for b in self.blocks if b.where == "body" and not b.cell]
        for position, block in enumerate(body):
            level = self.heading_level(block.element)
            title = paragraph_text(block.element).strip()
            if level is None or not title:
                continue
            below = []
            for following in body[position + 1:]:
                if self.heading_level(following.element) is not None:
                    break
                if block.element.getparent() is not following.element.getparent() or self._table_between(block.element, following.element):
                    break
                below.append(following)
            next_block = body[position + 1] if position + 1 < len(body) else None
            next_level = self.heading_level(next_block.element) if next_block else None
            if below and all(self._is_empty(b.element) for b in below):
                spot = Spot(below[0].id, "section")
            elif not below and (next_block is None or (next_level is not None and next_level <= level)) and not self._table_follows(block.element):
                spot = Spot(block.id, "after")
            else:
                continue
            item = Field("", title, "section", "section", f"Section under the heading \"{title}\"", "", [], [spot])
            found.append((self.blocks.index(block), 10**6, None, item))

        # One field for a placeholder that appears several times
        found.sort(key=lambda entry: (entry[0], entry[1]))
        fields: list[Field] = []
        groups: dict[str, Field] = {}
        for _, _, key, item in found:
            if key and key in groups:
                groups[key].spots.extend(item.spots)
                continue
            if not item.label:
                item.label = f"Blank {sum(1 for f in fields if f.how == item.how) + 1}"
            if key:
                groups[key] = item
            fields.append(item)
            if len(fields) >= MAX_FIELDS:
                break
        for number, item in enumerate(fields, 1):
            item.id = f"f{number}"
        return fields

    @staticmethod
    def _table_between(first: Any, second: Any) -> bool:
        node = first.getnext()
        while node is not None and node is not second:
            if node.tag == W_TBL:
                return True
            node = node.getnext()
        return False

    @staticmethod
    def _table_follows(paragraph: Any) -> bool:
        following = paragraph.getnext()
        return following is not None and following.tag == W_TBL

    # ── Writing ──

    def write_section(self, block: Block, mode: str, text: str) -> None:
        """Writes one paragraph per line under a heading, in the look of the empty paragraph that was there."""
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if not lines:
            return
        if mode == "after":
            anchor = OxmlElement("w:p")
            block.element.addnext(anchor)
        else:
            anchor = block.element
        pattern = empty_copy(anchor)
        bullet_style = self.style_id("List Bullet")
        for number, line in enumerate(lines):
            paragraph = anchor if number == 0 else copy.deepcopy(pattern)
            if number:
                anchor.addnext(paragraph)
                anchor = paragraph
            if _BULLET.match(line):
                line = _BULLET.sub("", line)
                if bullet_style:
                    properties = paragraph.find(W_PPR)
                    if properties is None:
                        properties = OxmlElement("w:pPr")
                        paragraph.insert(0, properties)
                    for old in properties.findall(qn("w:pStyle")):
                        properties.remove(old)
                    style = OxmlElement("w:pStyle")
                    style.set(W_VAL, bullet_style)
                    properties.insert(0, style)
                else:
                    line = "• " + line
            write_paragraph(paragraph, line)

    def fill(self, values: dict[str, str]) -> dict[str, Any]:
        """Writes one value per field. Fields without a value are left exactly as they are."""
        fields = self.fields()
        filled, merged, empty = [], [], []
        inline: dict[str, list[tuple[Spot, str, Field]]] = {}
        whole: list[tuple[Spot, str]] = []
        for item in fields:
            value = (values.get(item.id) or "").strip()
            if not value:
                empty.append(item.label)
                continue
            filled.append(item.label)
            for spot in item.spots:
                if spot.mode == "span":
                    inline.setdefault(spot.block, []).append((spot, value, item))
                else:
                    whole.append((spot, value))

        for block_id, items in inline.items():
            paragraph = self.block(block_id).element
            # Right to left, so the positions of the blanks still to come stay valid.
            for spot, value, item in sorted(items, key=lambda entry: entry[0].start, reverse=True):
                plain = item.how in ("blank", "sample")
                if replace_span(paragraph, spot.start, spot.end, value, plain=plain) and item.label not in merged:
                    merged.append(item.label)
        for spot, value in whole:
            block = self.block(spot.block)
            if spot.mode == "paragraph":
                write_paragraph(block.element, value)
            elif spot.mode == "append":
                append_text(block.element, value)
            else:
                self.write_section(block, spot.mode, value)
        return {"filled": filled, "empty": empty, "merged": merged}

    # ── Reading back ──

    def outline(self, limit: int = 220) -> str:
        """The file as numbered lines for a model: `p4 [Heading 1] Scope of work`."""
        lines = []
        for t, rows in enumerate(self.tables, 1):
            lines.append(f"(table {t} has {len(rows)} rows and {max((len(r) for r in rows), default=0)} columns)")
        for block in self.blocks:
            text = " ".join(paragraph_text(block.element).split())
            if not text and (not block.cell or block.where == "textbox"):
                continue
            notes = [self.style_name(block.element)] if block.where == "body" else [block.where.replace("textbox", "text box")]
            if block.cell:
                notes.append("table %d, row %d, column %d" % block.cell)
            align = self.alignment(block.element)
            if align != "left":
                notes.append(align)
            shown = text if len(text) <= limit else text[: limit - 3] + "..."
            lines.append(f"{block.id} [{', '.join(notes)}] {shown}")
        return "\n".join(lines)

    def preview(self) -> list[dict[str, Any]]:
        """The file as plain blocks for a rough on-screen preview: text, headings and tables, no fonts.

        Every paragraph keeps its id, also inside table cells, and a text box's
        paragraphs follow the paragraph the box hangs on.
        """
        boxes: dict[str, list[Block]] = {}
        for block in self.blocks:
            if block.anchor:
                boxes.setdefault(block.anchor, []).append(block)

        def entry(block: Block) -> dict[str, Any]:
            element = block.element
            return {
                "type": "p", "id": block.id, "text": paragraph_text(element), "level": self.heading_level(element),
                "align": self.alignment(element), "style": self.style_name(element), "box": block.where == "textbox",
            }

        def walk(parent: Any) -> list[dict[str, Any]]:
            items: list[dict[str, Any]] = []
            for child in parent.iterchildren():
                if child.tag == W_P:
                    block = self._by_element.get(id(child))
                    if block is None:
                        continue
                    inner = [entry(b) for b in boxes.get(block.id, []) if paragraph_text(b.element).strip()]
                    text = paragraph_text(child)
                    blank_before = not items or items[-1]["type"] != "p" or not items[-1]["text"]
                    if text.strip() or not (inner or blank_before):
                        items.append(entry(block))
                    items.extend(inner)
                elif child.tag == W_TBL:
                    rows = [[walk(cell) for cell in row.iterchildren(W_TC)] for row in child.iterchildren(W_TR)]
                    items.append({"type": "table", "rows": rows})
                elif child.tag == W_SDT:
                    content = child.find(qn("w:sdtContent"))
                    if content is not None:
                        items.extend(walk(content))
            return items

        return walk(self.doc.element.body)

    def save(self) -> bytes:
        self.sync_text_boxes()
        out = io.BytesIO()
        self.doc.save(out)
        data = out.getvalue()
        open_word(self.name, data)  # a file we cannot read back is not handed to anyone
        return data


def analyze(name: str, data: bytes) -> dict[str, Any]:
    word = WordFile(name, data)
    fields = word.fields()
    return {
        "name": name,
        "fields": [f.public() for f in fields],
        "preview": word.preview(),
        "paragraphs": len(word.blocks),
        "tables": len(word.tables),
        "letterhead": word.is_letterhead(),
    }


def fill(name: str, data: bytes, values: dict[str, str]) -> tuple[bytes, dict[str, Any]]:
    word = WordFile(name, data)
    report = word.fill(values)
    return word.save(), {**report, "preview": word.preview()}
