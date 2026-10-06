"""Changes a Word file by a fixed list of operations, and builds one from Markdown.

A model decides *which* operations to run; this file is what carries them
out. It only knows the operations listed in `OPERATIONS`. Anything else a
model asks for is skipped and reported, never guessed at, and every operation
is checked against the file before it is applied.

Paragraphs are addressed by the ids `WordFile` gives them (`p4`, `h1`, ...)
and tables by number, row and column, all counted from 1 - the same numbers
the model saw in the outline.
"""

from __future__ import annotations

import copy
import io
import re
from typing import Any

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .wordfile import (
    W_P, W_PPR, W_TBL, W_TC, W_TR, W_TXBX, Block, WordFile, clean_copy, empty_copy,
    paragraph_text, replace_span, text_runs, write_paragraph,
)

MAX_OPERATIONS = 60

# name -> what a model is told about it
OPERATIONS = {
    "replace_text": "Replace the whole text of one paragraph. Needs: ids (one id), text.",
    "insert": "Add new paragraph(s) before or after a paragraph, one per line of text. Needs: ids (one id), text. Optional: position (before/after), style.",
    "delete": "Remove paragraphs. Needs: ids.",
    "replace_section": "Replace everything under a heading, up to the next heading. Needs: ids (the heading's id), text (one paragraph per line; start a line with '- ' for a bullet).",
    "find_replace": "Replace a word or phrase everywhere in the document. Needs: find, replace. Optional: match_case.",
    "align": "Set alignment. Needs: value (left/center/right/justify) and either ids or target.",
    "format": "Set character formatting. Needs ids or target, and any of: bold, italic, underline (true/false), size (points), color (hex like 1F4E79), font (name).",
    "style": "Give paragraphs a named style of the document, such as 'Heading 1', 'Heading 2', 'Normal', 'Title', 'List Bullet'. Needs: ids, style.",
    "spacing": "Set line spacing. Needs: value (1, 1.15, 1.5 or 2) and either ids or target.",
    "set_cell": "Write into one table cell. Needs: table, row, col, text.",
    "add_row": "Add a row to a table, formatted like its last row. Needs: table, values (one per column). Optional: row (add after this row).",
    "delete_row": "Remove a table row. Needs: table, row.",
    "page": "Page layout. Any of: value (portrait/landscape), size (margins in centimetres).",
}
TARGETS = ("ids", "all", "headings", "body", "tables")

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT, "start": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER, "centre": WD_ALIGN_PARAGRAPH.CENTER, "centered": WD_ALIGN_PARAGRAPH.CENTER,
    "centred": WD_ALIGN_PARAGRAPH.CENTER, "middle": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT, "end": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "justified": WD_ALIGN_PARAGRAPH.JUSTIFY, "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_COLORS = {
    "black": "000000", "white": "FFFFFF", "red": "C00000", "blue": "1F4E79", "green": "2E7D32",
    "grey": "595959", "gray": "595959", "orange": "D9730D", "purple": "5B3A8C",
}
_STYLE_ALIASES = {
    "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "heading1": "Heading 1", "heading2": "Heading 2",
    "heading3": "Heading 3", "body": "Normal", "paragraph": "Normal", "bullet": "List Bullet", "bullets": "List Bullet",
    "list": "List Bullet", "number": "List Number", "numbered": "List Number",
}


class _Skip(Exception):
    """An operation that cannot be carried out; the message goes into the report."""


def _flag(value: Any) -> bool | None:
    if isinstance(value, bool) or value is None:
        return value
    return {"true": True, "yes": True, "1": True, "false": False, "no": False, "0": False}.get(str(value).strip().lower())


def _number(value: Any, what: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        raise _Skip(f"no {what} number given") from None


def _text(op: dict[str, Any]) -> str:
    text = op.get("text")
    if not isinstance(text, str) or not text.strip():
        raise _Skip("no text given")
    return text.strip()


class Editor:
    def __init__(self, word: WordFile):
        self.word = word

    # ── What an operation points at ──

    def _alive(self, block: Block) -> Block:
        if block.element.getparent() is None:
            raise _Skip(f"{block.id} was removed by an earlier change")
        return block

    def _blocks(self, op: dict[str, Any], *, one: bool = False) -> list[Block]:
        ids = op.get("ids") or []
        if isinstance(ids, str):
            ids = [ids]
        if op.get("id"):
            ids = [op["id"], *ids]
        ids = [str(i).strip() for i in ids if str(i).strip()]
        target = str(op.get("target") or ("ids" if ids else "")).lower()
        word = self.word
        if ids and target in ("", "ids"):
            try:
                blocks = [self._alive(word.block(i)) for i in ids]
            except KeyError as exc:
                raise _Skip(str(exc.args[0])) from None
        elif target == "all":
            blocks = [b for b in word.blocks if b.where in ("body", "textbox")]
        elif target == "headings":
            blocks = [b for b in word.blocks if word.heading_level(b.element) is not None or word.style_name(b.element) == "Title"]
        elif target == "body":
            blocks = [
                b for b in word.blocks
                if b.where == "body" and not b.cell and word.heading_level(b.element) is None and word.style_name(b.element) != "Title"
            ]
        elif target == "tables":
            blocks = [b for b in word.blocks if b.cell]
        else:
            raise _Skip("it does not say which paragraphs (ids) or which target (all, headings, body, tables)")
        blocks = [b for b in blocks if b.element.getparent() is not None]
        if not blocks:
            raise _Skip("nothing in the document matches")
        if one and len(blocks) != 1:
            raise _Skip("it needs exactly one paragraph id")
        return blocks

    def _cell(self, op: dict[str, Any]) -> Any:
        rows = self._table(op)
        r, c = _number(op.get("row"), "row"), _number(op.get("col") or op.get("column"), "column")
        if not 1 <= r <= len(rows) or not 1 <= c <= len(rows[r - 1]):
            raise _Skip(f"table {op.get('table')} has no row {r}, column {c}")
        return rows[r - 1][c - 1]

    def _table(self, op: dict[str, Any]) -> list[list[Any]]:
        t = _number(op.get("table") or 1, "table")
        if not 1 <= t <= len(self.word.tables):
            raise _Skip(f"there is no table {t}")
        return self.word.tables[t - 1]

    @staticmethod
    def _describe(blocks: list[Block]) -> str:
        if len(blocks) == 1:
            words = " ".join(paragraph_text(blocks[0].element).split())
            return f"\"{words[:40]}{'...' if len(words) > 40 else ''}\"" if words else blocks[0].id
        return f"{len(blocks)} paragraphs"

    # ── Building blocks ──

    def _pattern_for(self, anchor: Block, style: str | None) -> Any:
        """An empty paragraph in the look new text beside `anchor` should have."""
        word = self.word
        if style:
            style_id = word.style_id(_STYLE_ALIASES.get(style.lower().replace(" ", ""), style))
            if not style_id:
                raise _Skip(f"this document has no style named '{style}'")
            paragraph = OxmlElement("w:p")
            paragraph.style = style_id
            return paragraph
        if word.heading_level(anchor.element) is None:
            return empty_copy(anchor.element, neutral=True)
        # Next to a heading: body text, in the look of the first body paragraph that follows.
        node = anchor.element.getnext()
        while node is not None:
            if node.tag == W_P and word.heading_level(node) is None and paragraph_text(node).strip():
                return empty_copy(node, neutral=True)
            node = node.getnext()
        return OxmlElement("w:p")

    def _write_lines(self, after: Any, pattern: Any, text: str) -> int:
        """Adds one paragraph per line after `after`. Returns how many were added."""
        temporary = copy.deepcopy(pattern)
        after.addnext(temporary)
        self.word.write_section(Block("new", temporary, "body"), "section", text)
        return len([line for line in text.splitlines() if line.strip()])

    @staticmethod
    def _clear_cell(cell: Any, text: str) -> None:
        paragraphs = list(cell.iterchildren(W_P)) or list(cell.iter(W_P))
        if not paragraphs:
            paragraph = OxmlElement("w:p")
            cell.append(paragraph)
            paragraphs = [paragraph]
        write_paragraph(paragraphs[0], text)
        for extra in paragraphs[1:]:
            extra.getparent().remove(extra)

    # ── Operations ──

    def replace_text(self, op: dict[str, Any]) -> str:
        block = self._blocks(op, one=True)[0]
        was = self._describe([block])
        write_paragraph(block.element, _text(op))
        return f"Rewrote {was}"

    def insert(self, op: dict[str, Any]) -> str:
        block = self._blocks(op, one=True)[0]
        text = _text(op)
        pattern = self._pattern_for(block, op.get("style") if isinstance(op.get("style"), str) else None)
        before = str(op.get("position") or "after").lower() == "before"
        if before:
            # write after a marker placed in front of the anchor, then drop the marker
            marker = OxmlElement("w:p")
            block.element.addprevious(marker)
            count = self._write_lines(marker, pattern, text)
            marker.getparent().remove(marker)
        else:
            count = self._write_lines(block.element, pattern, text)
        return f"Added {count} paragraph{'s' if count != 1 else ''} {'before' if before else 'after'} {self._describe([block])}"

    def delete(self, op: dict[str, Any]) -> str:
        if str(op.get("target") or "").lower() in ("all", "body", "headings", "tables"):
            raise _Skip("deleting needs the ids of the paragraphs")
        blocks = self._blocks(op)
        removed = 0
        for block in blocks:
            element = block.element
            if element.find(f".//{qn('w:drawing')}") is not None or element.find(f".//{qn('w:pict')}") is not None:
                continue  # holds a picture
            parent = element.getparent()
            properties = element.find(W_PPR)
            holds_break = properties is not None and properties.find(qn("w:sectPr")) is not None
            only_one = parent.tag in (W_TC, W_TXBX) and len(list(parent.iterchildren(W_P))) == 1
            if holds_break or only_one:
                write_paragraph(element, "")  # the paragraph itself has to stay
            else:
                parent.remove(element)
            removed += 1
        if not removed:
            raise _Skip("those paragraphs hold pictures and were left alone")
        return f"Removed {removed} paragraph{'s' if removed != 1 else ''}"

    def replace_section(self, op: dict[str, Any]) -> str:
        block = self._blocks(op, one=True)[0]
        word = self.word
        if word.heading_level(block.element) is None:
            raise _Skip(f"{block.id} is not a heading")
        text = _text(op)
        under = []
        node = block.element.getnext()
        while node is not None and node.tag == W_P and word.heading_level(node) is None:
            if node.find(W_PPR) is not None and node.find(W_PPR).find(qn("w:sectPr")) is not None:
                break  # a section break ends the section
            under.append(node)
            node = node.getnext()
        filled = next((p for p in under if paragraph_text(p).strip()), None)
        pattern = empty_copy(filled if filled is not None else under[0], neutral=True) if under else OxmlElement("w:p")
        for paragraph in under:
            paragraph.getparent().remove(paragraph)
        count = self._write_lines(block.element, pattern, text)
        return f"Rewrote the section under {self._describe([block])} ({count} paragraph{'s' if count != 1 else ''})"

    def find_replace(self, op: dict[str, Any]) -> str:
        find, replacement = op.get("find"), op.get("replace")
        if not isinstance(find, str) or not find.strip() or not isinstance(replacement, str):
            raise _Skip("it needs the text to find and the text to put in its place")
        pattern = re.compile(re.escape(find), 0 if _flag(op.get("match_case")) else re.IGNORECASE)
        count = 0
        for block in self.word.blocks:
            element = block.element
            if element.getparent() is None or self.word.style_name(element).lower().startswith("toc"):
                continue
            matches = list(pattern.finditer(paragraph_text(element)))
            for match in reversed(matches):  # right to left keeps earlier positions valid
                replace_span(element, match.start(), match.end(), replacement)
            count += len(matches)
        if not count:
            raise _Skip(f"\"{find}\" does not appear in the document")
        return f"Replaced \"{find}\" with \"{replacement}\" in {count} place{'s' if count != 1 else ''}"

    def align(self, op: dict[str, Any]) -> str:
        value = _ALIGN.get(str(op.get("value") or "").strip().lower())
        if value is None:
            raise _Skip("alignment must be left, center, right or justify")
        blocks = self._blocks(op)
        for block in blocks:
            Paragraph(block.element, None).alignment = value
        return f"Aligned {self._describe(blocks)} {str(op['value']).strip().lower()}"

    def format(self, op: dict[str, Any]) -> str:
        bold, italic, underline = _flag(op.get("bold")), _flag(op.get("italic")), _flag(op.get("underline"))
        size, font = op.get("size"), op.get("font")
        color = str(op.get("color") or "").strip().lstrip("#")
        color = _COLORS.get(color.lower(), color.upper())
        if color and not re.fullmatch(r"[0-9A-F]{6}", color):
            raise _Skip(f"'{op.get('color')}' is not a colour it knows")
        if size is not None:
            try:
                size = float(size)
            except (TypeError, ValueError):
                raise _Skip("size must be a number of points") from None
            if not 6 <= size <= 72:
                raise _Skip("size must be between 6 and 72 points")
        if bold is None and italic is None and underline is None and not size and not color and not font:
            raise _Skip("it does not say what formatting to set")
        blocks = self._blocks(op)
        for block in blocks:
            for element in text_runs(block.element):
                style = Run(element, None).font
                if bold is not None:
                    style.bold = bold
                if italic is not None:
                    style.italic = italic
                if underline is not None:
                    style.underline = underline
                if size:
                    style.size = Pt(size)
                if color:
                    style.color.rgb = RGBColor.from_string(color)
                if isinstance(font, str) and font.strip():
                    style.name = font.strip()
        done = [name for name, on in (("bold", bold), ("italic", italic), ("underline", underline)) if on is not None]
        done += [f"{size:g} pt"] if size else []
        done += [f"colour {color}"] if color else []
        done += [font.strip()] if isinstance(font, str) and font.strip() else []
        return f"Formatted {self._describe(blocks)}: {', '.join(done)}"

    def style(self, op: dict[str, Any]) -> str:
        name = str(op.get("style") or op.get("value") or "").strip()
        name = _STYLE_ALIASES.get(name.lower().replace(" ", ""), name)
        style_id = self.word.style_id(name)
        if not style_id:
            raise _Skip(f"this document has no style named '{name}'")
        blocks = self._blocks(op)
        for block in blocks:
            block.element.style = style_id
        return f"Set {self._describe(blocks)} to the style {name}"

    def spacing(self, op: dict[str, Any]) -> str:
        try:
            value = float(op.get("value"))
        except (TypeError, ValueError):
            raise _Skip("line spacing must be a number such as 1, 1.5 or 2") from None
        if not 0.8 <= value <= 3:
            raise _Skip("line spacing must be between 0.8 and 3")
        blocks = self._blocks(op)
        for block in blocks:
            Paragraph(block.element, None).paragraph_format.line_spacing = value
        return f"Set line spacing of {self._describe(blocks)} to {value:g}"

    def set_cell(self, op: dict[str, Any]) -> str:
        cell = self._cell(op)
        if cell.getparent() is None:
            raise _Skip("that row was removed by an earlier change")
        text = op.get("text") if isinstance(op.get("text"), str) else ""
        self._clear_cell(cell, text.strip())
        return f"Set table {op.get('table') or 1}, row {op['row']}, column {op.get('col') or op.get('column')} to \"{text.strip()[:40]}\""

    def add_row(self, op: dict[str, Any]) -> str:
        rows = self._table(op)
        values = op.get("values") if isinstance(op.get("values"), list) else []
        live = [row for row in rows if row and row[0].getparent() is not None and row[0].getparent().getparent() is not None]
        if not live:
            raise _Skip("that table has no rows left")
        number = op.get("row")
        source = live[-1]
        if number is not None:
            index = _number(number, "row")
            if not 1 <= index <= len(rows) or rows[index - 1] not in live:
                raise _Skip(f"table {op.get('table') or 1} has no row {index}")
            source = rows[index - 1]
        original = source[0].getparent()
        row = copy.deepcopy(original)
        for paragraph in list(row.iter(W_P)):
            cleaned = clean_copy(paragraph)
            paragraph.getparent().replace(paragraph, cleaned)
        cells = list(row.iterchildren(W_TC))
        for position, cell in enumerate(cells):
            self._clear_cell(cell, str(values[position]).strip() if position < len(values) else "")
        original.addnext(row)
        rows.insert(rows.index(source) + 1, cells)
        return f"Added a row to table {op.get('table') or 1}"

    def delete_row(self, op: dict[str, Any]) -> str:
        rows = self._table(op)
        index = _number(op.get("row"), "row")
        if not 1 <= index <= len(rows):
            raise _Skip(f"table {op.get('table') or 1} has no row {index}")
        element = rows[index - 1][0].getparent()
        if element.getparent() is None:
            raise _Skip("that row was already removed")
        if len(list(element.getparent().iterchildren(W_TR))) <= 1:
            raise _Skip("a table has to keep at least one row")
        element.getparent().remove(element)
        return f"Removed row {index} of table {op.get('table') or 1}"

    def page(self, op: dict[str, Any]) -> str:
        done = []
        orientation = str(op.get("value") or op.get("orientation") or "").strip().lower()
        if orientation in ("portrait", "landscape"):
            for section in self.word.doc.sections:
                wide = section.page_width > section.page_height
                if wide != (orientation == "landscape"):
                    section.page_width, section.page_height = section.page_height, section.page_width
                section.orientation = WD_ORIENT.LANDSCAPE if orientation == "landscape" else WD_ORIENT.PORTRAIT
            done.append(orientation)
        elif orientation:
            raise _Skip("page orientation must be portrait or landscape")
        if op.get("size") is not None:
            try:
                margin = float(op["size"])
            except (TypeError, ValueError):
                raise _Skip("margins must be a number of centimetres") from None
            if not 0.5 <= margin <= 6:
                raise _Skip("margins must be between 0.5 and 6 cm")
            for section in self.word.doc.sections:
                section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(margin)
            done.append(f"{margin:g} cm margins")
        if not done:
            raise _Skip("it does not say what to change about the page")
        return f"Page set to {' and '.join(done)}"

    # ── Running a list ──

    def apply(self, operations: list[Any]) -> dict[str, list[str]]:
        applied, skipped = [], []
        for op in operations[:MAX_OPERATIONS]:
            if not isinstance(op, dict):
                skipped.append("One instruction was not understood")
                continue
            name = str(op.get("op") or op.get("action") or op.get("type") or "").strip().lower()
            if name not in OPERATIONS:
                skipped.append(f"'{name or '?'}' is not something it can do")
                continue
            try:
                applied.append(getattr(self, name)(op))
            except _Skip as reason:
                skipped.append(f"{name.replace('_', ' ')}: {reason}")
        if len(operations) > MAX_OPERATIONS:
            skipped.append(f"Only the first {MAX_OPERATIONS} changes were made")
        return {"applied": applied, "skipped": skipped}


def edit(name: str, data: bytes, operations: list[Any]) -> tuple[bytes, dict[str, Any]]:
    word = WordFile(name, data)
    report = Editor(word).apply(operations)
    result = word.save()
    return result, {**report, "preview": WordFile(name, result).preview()}


# ── A new document from Markdown ──

_INLINE = re.compile(r"(\*\*[^*\n]+\*\*|__[^_\n]+__|(?<![\w*])\*[^*\n]+\*(?![\w*])|(?<!\w)_[^_\n]+_(?!\w)|`[^`\n]+`)")
_LINK = re.compile(r"\[([^\]\n]+)\]\(([^)\s]+)\)")
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_BULLET = re.compile(r"^(\s*)[-*+•]\s+(.*)$")
_MD_NUMBER = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_MD_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_MD_DIVIDER = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _inline(paragraph: Any, text: str, *, bold: bool = False) -> None:
    """Adds text to a paragraph, turning **bold**, *italic* and `code` into real formatting."""
    text = _LINK.sub(r"\1 (\2)", text)
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith(("**", "__")) and len(piece) > 4:
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece.startswith(("*", "_")) and len(piece) > 2 and _INLINE.fullmatch(piece):
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(piece)
        if bold:
            run.bold = True


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def from_markdown(markdown: str, base: tuple[str, bytes] | None = None) -> bytes:
    """Builds a Word file from Markdown.

    With `base`, the new text replaces that file's body and takes its styles,
    page setup, headers and footers - a letterhead stays a letterhead.
    """
    lines = markdown.replace("\r\n", "\n").strip("\n").split("\n")
    if not any(line.strip() for line in lines):
        raise ValueError("there is nothing to write")
    if base:
        word = WordFile(*base)
        if not word.is_letterhead():
            # Its body is the design: replacing it would throw the layout away.
            raise ValueError("This file is a designed template. Use Fill a template so its layout is kept.")
        document = word.doc
        body = document.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        document = docx.Document()
    names = {style.name.lower(): style.name for style in document.styles}

    def styled(*wanted: str) -> Any:
        style = next((names[w.lower()] for w in wanted if w.lower() in names), None)
        return document.add_paragraph(style=style), style is not None

    # One "# Title" at the very top is the document's title, and the "##" under it are its
    # main headings: Title, Heading 1, Heading 2 - not Title, Heading 2, Heading 3.
    tops = [n for n, line in enumerate(lines) if re.match(r"^#\s+\S", line.strip())]
    opening = next((n for n, line in enumerate(lines) if line.strip()), 0)
    titled = tops == [opening]

    index, first, paragraph_open = 0, True, None
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(lines[index].replace("\t", "    ") or " ")
                run.font.name, run.font.size = "Consolas", Pt(9.5)
                index += 1
            index += 1
            paragraph_open, first = None, False
            continue

        if not stripped or _MD_RULE.match(line):
            paragraph_open = None
            index += 1
            continue

        heading = _MD_HEADING.match(stripped)
        if heading:
            depth = len(heading.group(1))
            is_title = titled and depth == 1
            level = 1 if is_title else min(depth - 1 if titled else depth, 3)
            wanted = ("Title",) if is_title else (f"Heading {level}",)
            paragraph, found = styled(*wanted, f"Heading {level}")
            _inline(paragraph, heading.group(2), bold=not found)
            if not found:
                for run in paragraph.runs:
                    run.font.size = Pt({1: 18, 2: 15, 3: 13}[level])
            paragraph_open, first = None, False
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and _MD_DIVIDER.match(lines[index + 1]) and "-" in lines[index + 1]:
            rows = [_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_cells(lines[index]))
                index += 1
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            if "table grid" in names:
                table.style = names["table grid"]
            else:
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    border = OxmlElement(f"w:{edge}")
                    for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "auto")):
                        border.set(qn(f"w:{key}"), value)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    _inline(table.rows[r].cells[c].paragraphs[0], value, bold=r == 0)
            paragraph_open, first = None, False
            continue

        bullet, number = _MD_BULLET.match(line), _MD_NUMBER.match(line)
        if bullet:
            deep = len(bullet.group(1).replace("\t", "    ")) >= 2
            paragraph, found = styled(*(("List Bullet 2",) if deep else ()), "List Bullet")
            _inline(paragraph, ("" if found else "• ") + bullet.group(2))
            if not found:
                paragraph.paragraph_format.left_indent = Cm(1.2 if deep else 0.6)
        elif number:
            # The number is written out: Word's automatic numbering would carry on from an earlier list.
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.75)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
            _inline(paragraph, f"{number.group(2)}.\t{number.group(3)}")
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            _inline(paragraph, stripped.lstrip("> ").strip())
            for run in paragraph.runs:
                run.italic = True
        elif paragraph_open is not None:
            # A single line break inside a paragraph stays a line break (addresses, sign-offs).
            paragraph_open.add_run().add_break()
            _inline(paragraph_open, stripped)
            index += 1
            continue
        else:
            paragraph_open = document.add_paragraph()
            _inline(paragraph_open, stripped)
            first = False
            index += 1
            continue
        paragraph_open, first = None, False
        index += 1

    out = io.BytesIO()
    document.save(out)
    data = out.getvalue()
    WordFile("document.docx", data)  # must read back
    return data
