"""Writes the reviewed rows as an Excel file or a CSV. No model here.

Both have exactly the columns of `checks.COLUMNS`, in that order, and a row for
each line of each invoice: an invoice of seven items is seven rows. As on the
page, the invoice's own values (number, date, seller, total, taxes,
confidence) stand on its first row; the rows after it carry the file name and
the line. Confidence is written as a percentage.

What can be asked for besides (`Options`):
  repeat_details  the invoice's number, date, seller, buyer ... on every row, for filtering
  repeat_totals   its total and taxes on every row too (a sum of the column then counts them again)
  checks          xlsx: a sheet listing what to check and why
  summary         xlsx: a sheet with one row for each invoice, no lines
  totals          xlsx: a last row adding up the amounts
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from .checks import AMOUNTS, COLUMNS, FIELDS, ITEM_FIELDS, KIND_NAMES, NUMBERS

WIDTHS = {
    "file": 34, "InvoiceId": 20, "InvoiceDate": 13, "DueDate": 13, "InvoiceTotal": 14, "VendorName": 30,
    "VendorAddress": 44, "CustomerName": 30, "CustomerId": 14, "BillingAddress": 44, "BillingAddressRecipient": 30,
    "VendorAddressRecipient": 30, "VendorGST": 19, "CustomerGST": 19, "confidence": 11,
    "Description": 44, "Qty": 12, "UnitPrice": 12, "UnitAmount": 13, "Discount": 12, "TaxableValue": 14,
    "CGSTAmount": 13, "SGSTAmount": 13, "IGSTAmount": 13, "TotalTaxAmount": 15, "HSNSAC": 12, "Currency": 9, "kind": 13,
    "lines": 8,
}
# The Summary sheet: one row for each invoice.
SUMMARY = (
    ("Invoice File Name", "file"), ("InvoiceId", "InvoiceId"), ("Invoice Date", "InvoiceDate"), ("VendorName", "VendorName"),
    ("VendorGST", "VendorGST"), ("CustomerName", "CustomerName"), ("Lines", "lines"), ("TaxableValue", "TaxableValue"),
    ("CGSTAmount", "CGSTAmount"), ("SGSTAmount", "SGSTAmount"), ("IGSTAmount", "IGSTAmount"),
    ("TotalTaxAmount", "TotalTaxAmount"), ("InvoiceTotal", "InvoiceTotal"), ("Currency", "Currency"),
    ("DocumentType", "kind"), ("Confidence", "confidence"),
)


@dataclass(frozen=True)
class Options:
    repeat_details: bool = False
    repeat_totals: bool = False
    checks: bool = False
    summary: bool = False
    totals: bool = False


def _lines(rows: list[dict[str, Any]]):
    """Every line of every invoice, as (invoice, line, whether it is the invoice's first)."""
    for row in rows:
        for number, item in enumerate(row.get("items") or [{}]):
            yield row, item, number == 0


def _cell(row: dict[str, Any], item: dict[str, Any], first: bool, key: str, options: Options) -> Any:
    if key == "file":
        return row.get("file") or ""
    if key in ITEM_FIELDS:
        return (item.get("values") or {}).get(key) or ""
    if key == "lines":
        return len([i for i in row.get("items") or [] if any((i.get("values") or {}).values())]) or ""
    if not first and not (options.repeat_totals if key in AMOUNTS else options.repeat_details):
        return ""  # as on the page: the invoice's own values stand on its first row
    if key == "confidence":
        return row.get("confidence")
    if key == "kind":
        return KIND_NAMES.get(row.get("kind") or "", "")
    return (row.get("values") or {}).get(key) or ""


def _level(row: dict[str, Any], item: dict[str, Any], first: bool, key: str) -> str:
    if key in ITEM_FIELDS:
        return (item.get("levels") or {}).get(key, "")
    return (row.get("levels") or {}).get(key, "") if first else ""  # shaded once, on the invoice's first row


def _number(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _percent(value: Any) -> str:
    number = _number(value)
    return "" if number is None else f"{round(number * 100)}%"


def _safe(text: str) -> str:
    """A cell starting with = or @ would be run by Excel as a formula when a CSV is opened.
    Invoices come from outside, so such a value gets a leading apostrophe."""
    risky = text[:1] in ("=", "@", "\t", "\r") or (text[:1] in ("+", "-") and _number(text) is None)
    return "'" + text if risky else text


def to_csv(rows: list[dict[str, Any]], options: Options = Options()) -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow([title for title, _ in COLUMNS])
    for row, item, first in _lines(rows):
        line = []
        for _, key in COLUMNS:
            value = _cell(row, item, first, key, options)
            line.append(_percent(value) if key == "confidence" else _safe(str(value if value is not None else "")))
        writer.writerow(line)
    return out.getvalue().encode("utf-8-sig")  # the BOM makes Excel read ₹, accents and other scripts correctly


def to_xlsx(rows: list[dict[str, Any]], date_format: str, options: Options = Options()) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    head = Font(bold=True)
    doubtful = PatternFill("solid", fgColor="FFF2CC")

    def put(cell, key: str, value: Any) -> None:
        if key in ("InvoiceDate", "DueDate") and value:
            try:
                cell.value, cell.number_format = datetime.strptime(value, date_format), "DD-MMM-YYYY"
            except ValueError:
                cell.value, cell.data_type = value, "s"
        elif key in NUMBERS and _number(value) is not None:
            cell.value, cell.number_format = _number(value), "#,##0.00"
        elif key == "confidence":
            cell.value, cell.number_format = _number(value), "0%"  # 0.8 is shown as 80%
        elif key == "lines":
            cell.value = value or None
        elif value != "":
            cell.value = str(value)
            cell.data_type = "s"  # always text: never a formula, and an id like 00123 keeps its zeros
        cell.alignment = Alignment(wrap_text=key in ("VendorAddress", "BillingAddress", "Description"), vertical="top")

    def fill(sheet, columns, lines, options: Options) -> dict[int, int]:
        sheet.append([title for title, _ in columns])
        for index, (_, key) in enumerate(columns, 1):
            sheet.cell(1, index).font = head
            sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[key]
        starts: dict[int, int] = {}  # where each invoice's first row is
        sums: dict[int, float] = {}
        number = 1
        for number, (row, item, first) in enumerate(lines, 2):
            if first:
                starts[id(row)] = number
            for index, (_, key) in enumerate(columns, 1):
                value, cell = _cell(row, item, first, key, options), sheet.cell(number, index)
                put(cell, key, value)
                if _level(row, item, first, key) == "check":
                    cell.fill = doubtful
                # An invoice's figures are added once, whatever is repeated on its other rows.
                if key in NUMBERS and key != "UnitPrice" and (first or key in ITEM_FIELDS) and _number(value) is not None:
                    sums[index] = sums.get(index, 0.0) + _number(value)
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        if options.totals and sums:
            last = number + 2  # a gap, so the filter and a sort leave the totals row alone
            sheet.cell(last, 1).value, sheet.cell(last, 1).font = "Total", head
            for index, total in sums.items():
                cell = sheet.cell(last, index)
                cell.value, cell.number_format, cell.font = round(total, 2), "#,##0.00", head
        return starts

    sheet = book.active
    sheet.title = "Invoices"
    starts = fill(sheet, COLUMNS, _lines(rows), options)

    if options.summary:
        fill(book.create_sheet("Summary"), SUMMARY, ((row, {}, True) for row in rows), options)

    if options.checks:
        checks = book.create_sheet("Checks")
        checks.append(["Invoice File Name", "Row", "Pages", "Kind", "Fields", "Level", "Note"])
        for index, width in enumerate((34, 6, 8, 10, 34, 8, 90), 1):
            checks.cell(1, index).font = head
            checks.column_dimensions[get_column_letter(index)].width = width
        for row in rows:
            pages = ", ".join(str(p) for p in row.get("pages") or [])
            levels, notes = row.get("levels") or {}, row.get("notes") or {}
            # What needs a look, and only that: a value marked for checking, or a note on a value
            # that nothing confirmed. A field that is simply not on the invoice is not listed.
            # Fields that share a note (the three taxes of a sum that is off) are one line.
            found: dict[tuple[str, str], list[str]] = {}
            for field in (*FIELDS, *ITEM_FIELDS):
                level, note = levels.get(field, ""), notes.get(field, "")
                if level == "check" or (note and level == "read"):
                    found.setdefault((level, note), []).append(field)
            for (level, note), fields in found.items():
                checks.append([row.get("file") or "", starts.get(id(row)), pages, row.get("kind") or "", ", ".join(fields), level, note])
        checks.freeze_panes = "A2"

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
