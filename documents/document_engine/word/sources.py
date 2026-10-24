"""A source file -> the text the model reads. No model here.

PDF, Word and plain text go through the chat's own reader (`extract_text`).
Excel is read here: each sheet becomes a Markdown table, so the model sees
which value belongs to which column.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import PurePath

from ollama_pipeline import UnsupportedDocument, extract_text

ACCEPTED = (".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md", ".json")
MAX_CHARS = 200_000  # the same limit the chat puts on an attached file
MAX_ROWS = 2000      # rows read from one sheet

_PICTURES = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".gif")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("|", "/").replace("\n", " ").strip()


def _excel(name: str, data: bytes) -> str:
    from openpyxl import load_workbook  # here, so the app starts without it

    try:
        # data_only: the values Excel last worked out, not the formulas.
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedDocument(f"{name} could not be read as an Excel file") from exc
    parts = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(cells)
            if len(rows) > MAX_ROWS:
                break
        if not rows:
            continue
        width = max(len(r) for r in rows)
        lines = ["| " + " | ".join(r + [""] * (width - len(r))) + " |" for r in rows[:MAX_ROWS]]
        lines.insert(1, "|" + " --- |" * width)
        if len(rows) > MAX_ROWS:
            lines.append(f"(only the first {MAX_ROWS} rows of this sheet were read)")
        parts.append(f"Sheet: {sheet.title}\n" + "\n".join(lines))
    book.close()
    return "\n\n".join(parts)


def read_source(name: str, data: bytes, limit: int = MAX_CHARS) -> tuple[str, bool]:
    """Returns (text, truncated). Raises UnsupportedDocument with a reason the user can act on."""
    suffix = PurePath(name).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        text = _excel(name, data).strip()
        if not text:
            raise UnsupportedDocument(f"{name} has no values in it")
        return text[:limit], len(text) > limit
    if suffix == ".xls":
        raise UnsupportedDocument(f"{name} is an old .xls file. Save it as .xlsx and add it again.")
    if suffix == ".doc":
        raise UnsupportedDocument(f"{name} is an old .doc file. Save it as .docx and add it again.")
    if suffix in _PICTURES:
        raise UnsupportedDocument(f"{name} is a picture. This tab reads text: PDF, Word, Excel, CSV and text files.")
    return extract_text(name, data, limit)
