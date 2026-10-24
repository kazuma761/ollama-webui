"""A spreadsheet (Excel or CSV) -> SQLite tables the model can query. No model here.

A big sheet does not fit in a model's context, and a model that reads 20,000
rows of numbers adds them up wrong anyway. Loaded into SQLite, the model asks
for what it needs with SQL (`SELECT region, SUM(amount) ... GROUP BY region`)
and the database does the counting: the answer is exact and only a few rows
reach the model. AnythingLLM turns a sheet into text and embeds it instead,
which is fine for finding a row but not for totals.

Each sheet becomes one table. The header row is found even under title rows
("Sales report, Q1"), column names become plain SQL names (the original
header is kept beside each), and each column gets a type: numbers stay
numbers ("1,23,456.50" and "₹500" included), dates become YYYY-MM-DD.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import sqlite3
from pathlib import Path, PurePath
from typing import Any

SPREADSHEETS = (".xlsx", ".xlsm", ".csv", ".tsv")
MAX_ROWS = 1_000_000  # rows read from one sheet
SAMPLE = 5  # rows shown to the model as an example
SCAN = 30  # rows looked at to find the header
BATCH = 5000

_NUMBER = re.compile(r"^\s*[-+]?(?:₹|rs\.?|inr|\$|€|£)?\s*[-+]?(?:\d{1,3}(?:,\d{2,3})+|\d+)(?:\.\d+)?\s*%?\s*$", re.I)
_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%b %d, %Y",
                 "%d %B %Y", "%B %d, %Y", "%Y/%m/%d", "%d/%m/%y")


class SheetError(ValueError):
    """The file cannot be read as a spreadsheet; the message says why."""


def is_spreadsheet(name: str) -> bool:
    return PurePath(name).suffix.lower() in SPREADSHEETS


def _name(text: str, taken: set[str], fallback: str) -> str:
    """A plain SQL name (letters, digits, _), not taken yet."""
    base = re.sub(r"[^a-z0-9]+", "_", str(text).lower()).strip("_")[:48] or fallback
    if base[0].isdigit():
        base = "c_" + base
    name, n = base, 2
    while name in taken:
        name, n = f"{base}_{n}", n + 1
    taken.add(name)
    return name


def _number(value: str) -> float | None:
    if not _NUMBER.match(value):
        return None
    cleaned = re.sub(r"[^\d.\-+]", "", value.replace(",", ""))
    try:
        return float(cleaned)
    except ValueError:
        return None


def _date(value: str) -> str | None:
    value = value.strip()
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _cell(value: Any) -> Any:
    """An Excel cell as Python: numbers stay numbers, dates become ISO text, blanks None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() and abs(value) < 2**53 else value
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _header_at(rows: list[list[Any]]) -> int | None:
    """Which of the first rows is the header: the first wide row that is mostly words."""
    if not rows:
        return None
    widest = max(sum(v is not None for v in row) for row in rows)
    for i, row in enumerate(rows):
        filled = [v for v in row if v is not None]
        if len(filled) >= max(2, widest * 0.6) or (widest == 1 and filled):
            words = sum(isinstance(v, str) and _number(v) is None and _date(v) is None for v in filled)
            return i if words >= len(filled) * 0.6 else None
    return None


def _column_type(values: list[Any]) -> str:
    """INTEGER, REAL, DATE or TEXT, from a column's values (most of them must agree)."""
    seen = [v for v in values if v is not None]
    if not seen:
        return "TEXT"
    numbers = [v for v in seen if isinstance(v, (int, float))]
    texts = [v for v in seen if isinstance(v, str)]
    if any(re.match(r"^0\d", v) for v in texts):
        return "TEXT"  # codes with leading zeros (PIN codes, 00123) stay as written
    parsed = [n for n in (_number(v) for v in texts) if n is not None]
    if numbers or parsed:
        if len(numbers) + len(parsed) >= 0.95 * len(seen):
            whole = all(isinstance(v, int) for v in numbers) and all(n.is_integer() for n in parsed)
            return "INTEGER" if whole else "REAL"
    if texts and sum(_date(v) is not None for v in texts) >= 0.95 * len(seen):
        return "DATE"
    return "TEXT"


def _convert(value: Any, kind: str) -> Any:
    if value is None:
        return None
    if kind in ("REAL", "INTEGER") and isinstance(value, str):
        n = _number(value)
        return value if n is None else (int(n) if n.is_integer() and kind == "INTEGER" else n)
    if kind == "DATE" and isinstance(value, str):
        return _date(value) or value
    return value


def _sheets(name: str, data: bytes):
    """(sheet name, an iterator of rows as lists of Python values) for each sheet."""
    suffix = PurePath(name).suffix.lower()
    if suffix in (".csv", ".tsv"):
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|") if suffix == ".csv" else csv.excel_tab
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(io.StringIO(text), dialect)
        yield PurePath(name).stem, ([_cell(v) for v in row] for row in rows)
        return
    try:
        from openpyxl import load_workbook  # here, so the app starts without it
    except ImportError as exc:
        raise SheetError("Reading Excel files needs the openpyxl package on the server: pip install -r requirements.txt") from exc
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise SheetError(f"{name} could not be read as an Excel file.") from exc
    try:
        for sheet in book.worksheets:
            yield sheet.title, ([_cell(v) for v in row] for row in sheet.iter_rows(values_only=True))
    finally:
        book.close()


def import_file(name: str, data: bytes, path: Path, max_rows: int = MAX_ROWS) -> list[dict[str, Any]]:
    """Loads every sheet of the file into the SQLite file at `path`. Returns what each table holds."""
    if not is_spreadsheet(name):
        raise SheetError(f"{name} is not a spreadsheet (Excel .xlsx or CSV).")
    tables: list[dict[str, Any]] = []
    taken: set[str] = set()
    db = sqlite3.connect(path)
    try:
        for sheet, rows in _sheets(name, data):
            table = _load(db, sheet, rows, taken, max_rows)
            if table:
                tables.append(table)
        db.commit()
    finally:
        db.close()
    if not tables:
        raise SheetError(f"{name} has no rows with values in it.")
    return tables


TYPE_ROWS = 1000  # rows looked at to decide a column's type


def _load(db: sqlite3.Connection, sheet: str, rows, taken: set[str], max_rows: int) -> dict[str, Any] | None:
    """One sheet into one table; None for an empty sheet."""
    first: list[list[Any]] = []
    for row in rows:
        if any(v is not None for v in row):
            first.append(row)
            if len(first) >= SCAN:
                break
    if not first:
        return None
    at = _header_at(first)
    used = max(max((i for i, v in enumerate(r) if v is not None), default=-1) for r in first) + 1
    header = ((first[at] if at is not None else []) + [None] * used)[:used]
    title = " / ".join(" ".join(str(v) for v in r if v is not None) for r in first[: at or 0])[:200]

    # More rows, to decide each column's type on enough of them.
    buffered = first[at + 1:] if at is not None else first
    for row in rows:
        if any(v is not None for v in row):
            buffered.append(row)
            if len(buffered) >= TYPE_ROWS:
                break
    cells = [(r + [None] * used)[:used] for r in buffered]

    names: set[str] = set()
    columns = []
    for i in range(used):
        label = header[i]
        columns.append({
            "name": _name(label if label is not None else f"column_{i + 1}", names, f"column_{i + 1}"),
            "header": "" if label is None else str(label),
            "type": _column_type([r[i] for r in cells]),
        })

    table = _name(sheet, taken, "sheet")
    definitions = ", ".join('"%s" %s' % (c["name"], "TEXT" if c["type"] == "DATE" else c["type"]) for c in columns)
    db.execute(f'CREATE TABLE "{table}" ({definitions})')
    insert = f'INSERT INTO "{table}" VALUES ({", ".join("?" * len(columns))})'
    kinds = [c["type"] for c in columns]
    count, batch, sample, truncated = 0, [], [], False

    def take(row: list[Any]) -> None:
        nonlocal count
        values = [_convert(v, k) for v, k in zip((row + [None] * used)[:used], kinds)]
        if any(v is not None for v in values):
            batch.append(values)
            if len(sample) < SAMPLE:
                sample.append(values)
            count += 1
            if len(batch) >= BATCH:
                db.executemany(insert, batch)
                batch.clear()

    for row in cells:
        take(row)
    for row in rows:
        if count >= max_rows:
            truncated = True
            break
        take(row)
    if batch:
        db.executemany(insert, batch)
    return {"table": table, "sheet": sheet, "rows": count, "columns": columns, "sample": sample,
            "title": title, "truncated": truncated}


def describe(tables: list[dict[str, Any]], file: str) -> str:
    """What the model is told about a loaded file: its tables, their columns and a few rows."""
    out = []
    for t in tables:
        lines = [f'Table "{t["table"]}" (sheet "{t["sheet"]}" of {file}, {t["rows"]:,} rows'
                 + (", cut at the row limit" if t["truncated"] else "") + ")"]
        if t["title"]:
            lines.append(f"Title above the table: {t['title']}")
        cols = ", ".join(f'{c["name"]} {c["type"]}' + (f' ("{c["header"]}")' if c["header"] and c["header"] != c["name"] else "")
                         for c in t["columns"])
        lines.append(f"Columns: {cols}")
        if t["sample"]:
            names = [c["name"] for c in t["columns"]]
            lines.append("First rows:")
            lines.append("| " + " | ".join(names) + " |")
            for row in t["sample"]:
                lines.append("| " + " | ".join("" if v is None else str(v)[:40] for v in row) + " |")
        out.append("\n".join(lines))
    return "\n\n".join(out)


def as_markdown(path: Path, tables: list[dict[str, Any]], max_rows: int = 400) -> str:
    """A small sheet, whole, as Markdown tables: for a file that fits in the model's context as it is."""
    db = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        out = []
        for t in tables:
            names = [c["name"] for c in t["columns"]]
            rows = db.execute(f'SELECT * FROM "{t["table"]}" LIMIT {int(max_rows)}').fetchall()
            lines = [f'Sheet "{t["sheet"]}" ({t["rows"]:,} rows)', "| " + " | ".join(names) + " |",
                     "|" + " --- |" * len(names)]
            lines += ["| " + " | ".join("" if v is None else str(v) for v in row) + " |" for row in rows]
            out.append("\n".join(lines))
        return "\n\n".join(out)
    finally:
        db.close()
